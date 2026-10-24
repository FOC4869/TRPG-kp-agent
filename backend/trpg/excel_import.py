"""从「秋叶EXODUS」COC 角色卡 Excel 导入角色。

只依赖 openpyxl（惰性导入），读取「人物卡」sheet 的固定单元格布局：
- 属性：力量 U3 / 敏捷 AA3 / 意志 AG3 / 体质 U5 / 外貌 AA5 / 教育 AG5 / 体型 U7 / 智力 AA7 / 幸运 AG7
- 派生：体力 E10(现)/G10(上限)、理智 N10(现)/P10(上限)、魔法 W10(现)/Y10(上限)、移动力 AF10、
        伤害加值 AP52、体格 AP55
- 技能：两列表格（左列 名称 F / 子名 H / 成功率 R；右列 名称 AB / 成功率 AN；行 16~49）
"""

from __future__ import annotations

import re
from typing import Optional

from .character import Character

_SUB_MARK = re.compile(r"[①②③④]")
_SKIP_NAMES = {"——", "自定义技能", ""}


def load_character_from_xlsx(path: str) -> Character:
    import openpyxl  # 惰性导入，核心包不依赖 openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["人物卡"]

    name = str(_cell(ws, "E3") or "")
    age = _int(_cell(ws, "E4"))
    occupation = str(_cell(ws, "E5") or "")

    attributes = {
        "STR": _int(_cell(ws, "U3")),
        "DEX": _int(_cell(ws, "AA3")),
        "POW": _int(_cell(ws, "AG3")),
        "CON": _int(_cell(ws, "U5")),
        "APP": _int(_cell(ws, "AA5")),
        "EDU": _int(_cell(ws, "AG5")),
        "SIZ": _int(_cell(ws, "U7")),
        "INT": _int(_cell(ws, "AA7")),
    }
    luck = _int(_cell(ws, "AG7"))

    hp = _int(_cell(ws, "E10"))
    hp_max = _int(_cell(ws, "G10"))
    san = _int(_cell(ws, "N10"))
    san_max = _int(_cell(ws, "P10"))
    mp = _int(_cell(ws, "W10"))
    mp_max = _int(_cell(ws, "Y10"))
    mov = _int(_cell(ws, "AF10"))
    db = str(_cell(ws, "AP52") or "0").strip()
    build = _int(_cell(ws, "AP55"))

    skills = _parse_skills(ws)

    return Character(
        name=name, age=age, occupation=occupation,
        attributes=attributes, skills=skills,
        hp=hp, hp_max=hp_max, san=san, san_max=san_max,
        mp=mp, mp_max=mp_max, luck=luck, mov=mov, db=db, build=build,
    )


def _cell(ws, coord):
    v = ws[coord].value
    return None if v is None else v


def _int(v) -> int:
    if v is None:
        return 0
    try:
        return int(v)
    except (TypeError, ValueError):
        return 0


def _parse_skills(ws) -> dict:
    skills: dict = {}
    for row in range(16, 50):
        # 左列：名称 F、子名 H、成功率 R
        _add_skill(skills, ws[f"F{row}"].value, ws[f"H{row}"].value, ws[f"R{row}"].value)
        # 右列：名称 AB、成功率 AN
        _add_skill(skills, ws[f"AB{row}"].value, None, ws[f"AN{row}"].value)
    return skills


def _add_skill(skills: dict, raw_name, sub_name, raw_value) -> None:
    name = _clean_skill_name(raw_name, sub_name)
    if name is None or raw_value is None:
        return
    # 同名（如外语①②③ 聚合）取最大值
    skills[name] = max(skills.get(name, 0), _int(raw_value))


def _clean_skill_name(raw_name, sub_name) -> Optional[str]:
    if raw_name is None:
        return None
    name = str(raw_name).strip()
    if not name or name in _SKIP_NAMES:
        return None
    name = name.replace("Ω", "").strip()
    # 主技能带子技能名（格斗：/射击：）→ 用子技能名（斗殴/手枪）
    if name.rstrip("：") in ("格斗", "射击") and sub_name and str(sub_name).strip():
        return str(sub_name).strip()
    # 去掉 ①②③ 标记（外语①→外语）
    name = _SUB_MARK.sub("", name)
    name = name.rstrip("：").strip()
    return name or None
